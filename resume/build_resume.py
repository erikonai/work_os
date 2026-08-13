#!/usr/bin/env python3
"""Build Erik Leavell's two-column resumes (dark sidebar + main column) as .docx.

Two variants:
  general  -> finance-executive framing
  abacum   -> applied-AI framing for Abacum roles

Layout is a single full-bleed 1x2 table: shaded sidebar cell, white main cell.
Row height is pinned so the dark sidebar paints the full page.
"""

import copy
import sys
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Emu

import content

FONT = "Liberation Sans"

DARK = "1B2735"          # sidebar background
ACCENT = "7FB3D9"        # sidebar section headings
SIDE_TEXT = "D9E2EC"     # sidebar body text
INK = "1B2735"           # main column headings
BODY = "2E3A46"          # main column body
MUTED = "5B6B7B"         # dates / locations
RULE = "C3CDD7"

PAGE_W = Inches(8.5)
PAGE_H = Inches(11.0)
SIDE_W = Inches(2.42)
MAIN_W = Inches(6.08)
MAIN_PAD_L, MAIN_PAD_R = 0.30, 0.34
MAIN_TEXT_W = 6.08 - MAIN_PAD_L - MAIN_PAD_R  # right tab stop for dates


# ---------------------------------------------------------------- xml helpers

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def shade(cell, hex_fill):
    cell._tc.get_or_add_tcPr().append(
        _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": hex_fill}))


def cell_margins(cell, top, left, bottom, right):
    mar = _el("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        mar.append(_el(f"w:{tag}", **{"w:w": int(val * 1440), "w:type": "dxa"}))
    cell._tc.get_or_add_tcPr().append(mar)


def configure_table(table, width_in):
    """Rebuild tblPr with children in the order the OOXML schema requires."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)

    tblPr.append(_el("w:tblW", **{"w:w": int(width_in * 1440), "w:type": "dxa"}))
    tblPr.append(_el("w:tblInd", **{"w:w": 0, "w:type": "dxa"}))

    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", **{"w:val": "none", "w:sz": 0, "w:space": 0,
                                           "w:color": "auto"}))
    tblPr.append(borders)

    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))

    mar = _el("w:tblCellMar")
    for tag in ("top", "start", "bottom", "end"):
        mar.append(_el(f"w:{tag}", **{"w:w": 0, "w:type": "dxa"}))
    tblPr.append(mar)


def set_grid(table, widths_in):
    """Under fixed layout the renderer uses w:tblGrid, not w:tcW. python-docx
    leaves the grid at equal columns, so set it explicitly or the split is 50/50."""
    grid = table._tbl.find(qn("w:tblGrid"))
    for child in list(grid):
        grid.remove(child)
    for w in widths_in:
        grid.append(_el("w:gridCol", **{"w:w": int(w * 1440)}))


def bottom_rule(par, hex_color=RULE, size=6):
    pbdr = _el("w:pBdr")
    pbdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": size,
                                   "w:space": 2, "w:color": hex_color}))
    par._p.get_or_add_pPr().append(pbdr)


def right_tab(par, pos_in):
    """Right-aligned tab via python-docx so w:tabs lands in schema order in pPr.
    Appending it manually puts it after w:spacing/w:ind, which LibreOffice drops."""
    par.paragraph_format.tab_stops.add_tab_stop(
        Inches(pos_in), WD_TAB_ALIGNMENT.RIGHT)


def letter_space(run, twentieths):
    run._r.get_or_add_rPr().append(_el("w:spacing", **{"w:val": twentieths}))


def hyperlink(par, url, text, size, color, bold=False, italic=False, underline=False):
    r_id = par.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = _el("w:hyperlink", **{"r:id": r_id})
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = _el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT})
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if underline:
        rpr.append(_el("w:u", **{"w:val": "single"}))
    rpr.append(_el("w:color", **{"w:val": color}))
    rpr.append(_el("w:sz", **{"w:val": int(size * 2)}))
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


# ------------------------------------------------------------- text primitives

def para(cell, space_before=0, space_after=0, line=1.0, align=None, first=False):
    p = cell.paragraphs[0] if (first and cell.paragraphs) else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.widow_control = False
    if align:
        p.alignment = align
    return p


def text(par, s, size, color, bold=False, italic=False, caps=False, spacing=None):
    r = par.add_run(s.upper() if caps else s)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    r._r.get_or_add_rPr().append(
        _el("w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT, "w:cs": FONT}))
    if spacing:
        letter_space(r, spacing)
    return r


def bullet(cell, s_bold, s_rest, size, color, bold_color=None, space_before=0,
           space_after=2.4, line=1.0, indent=0.13):
    """Hanging-indent bullet rendered with a literal glyph (keeps spacing tight)."""
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.left_indent = Inches(indent)
    pf.first_line_indent = Inches(-indent)
    pf.widow_control = False
    text(p, "– ", size, color)
    if s_bold:
        text(p, s_bold, size, bold_color or color, bold=True)
    if s_rest:
        text(p, s_rest, size, color)
    return p


# ------------------------------------------------------------------ page setup

def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = PAGE_W, PAGE_H
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0)
    s.header_distance = s.footer_distance = Inches(0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    return doc


def build_frame(doc, sidebar_pad, main_pad):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False
    configure_table(table, 8.5)
    set_grid(table, [2.42, 6.08])

    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    row.height = Inches(11.0)

    side, main = row.cells
    side.width, main.width = SIDE_W, MAIN_W
    shade(side, DARK)
    cell_margins(side, *sidebar_pad)
    cell_margins(main, *main_pad)
    return side, main


# --------------------------------------------------------------- section heads

def side_head(cell, label, sz, gap_before):
    p = para(cell, space_before=gap_before, space_after=2.6)
    text(p, label, sz, ACCENT, bold=True, caps=True, spacing=24)
    bottom_rule(p, "3E5468", 4)


def main_head(cell, label, sz, gap_before, first=False):
    p = para(cell, space_before=gap_before, space_after=3.4, first=first)
    text(p, label, sz, INK, bold=True, caps=True, spacing=30)
    bottom_rule(p, RULE, 6)


# ---------------------------------------------------------------- the two docs

def build(variant, out_path, t):
    """t = typography/spacing dial, tuned per variant to land on one page."""
    c = content.GENERAL if variant == "general" else content.ABACUM
    doc = new_doc()
    side, main = build_frame(
        doc,
        sidebar_pad=(t["side_top"], 0.26, 0.10, 0.24),
        main_pad=(t["main_top"], MAIN_PAD_L, 0.10, MAIN_PAD_R),
    )

    # ---- sidebar: identity block
    p = para(side, space_after=0, line=0.86, first=True)
    text(p, "ERIK", t["name"], "FFFFFF", bold=True, spacing=40)
    p = para(side, space_after=t["name_gap"], line=0.86)
    text(p, "LEAVELL", t["name"], ACCENT, bold=True, spacing=40)

    for line in content.CONTACT:
        p = para(side, space_after=1.2, line=1.0)
        text(p, line, t["contact"], SIDE_TEXT)

    # ---- sidebar: blocks
    for block in c["sidebar"]:
        side_head(side, block["title"], t["side_head"], t["side_gap"])
        for item in block["items"]:
            if isinstance(item, tuple):
                lead, rest = item
                p = para(side, space_after=t["side_item_gap"], line=t["side_line"])
                text(p, lead, t["side_body"], "FFFFFF", bold=True)
                if rest:
                    text(p, " " + rest, t["side_body"], SIDE_TEXT)
            else:
                p = para(side, space_after=t["side_item_gap"], line=t["side_line"])
                text(p, item, t["side_body"], SIDE_TEXT)

    # ---- main: summary
    main_head(main, c["summary_head"], t["main_head"], 0, first=True)
    p = para(main, space_after=0, line=t["body_line"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    text(p, c["summary"], t["body"], BODY)

    # ---- main: AI systems (abacum only)
    if c.get("ai_systems"):
        main_head(main, "AI SYSTEMS I'VE BUILT", t["main_head"], t["main_gap"])
        for sysm in c["ai_systems"]:
            p = para(main, space_after=1.6, line=1.0)
            if sysm.get("url"):
                hyperlink(p, sysm["url"], sysm["name"], t["role"], INK, bold=True)
            else:
                text(p, sysm["name"], t["role"], INK, bold=True)
            text(p, "  ·  " + sysm["tag"], t["body"], MUTED, italic=True)
            p = para(main, space_after=t["sys_gap"], line=t["body_line"],
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            text(p, sysm["desc"], t["body"], BODY)

    # ---- main: experience
    main_head(main, "EXPERIENCE", t["main_head"], t["main_gap"])
    for i, job in enumerate(c["experience"]):
        p = para(main, space_before=0 if i == 0 else t["job_gap"],
                 space_after=1.0, line=1.0)
        right_tab(p, MAIN_TEXT_W)
        text(p, job["title"], t["role"], INK, bold=True)
        text(p, "  ·  " + job["company"], t["role"], BODY)
        text(p, "\t" + job["dates"], t["meta"], MUTED)
        for b in job["bullets"]:
            bullet(main, b[0] if isinstance(b, tuple) else None,
                   b[1] if isinstance(b, tuple) else b,
                   t["body"], BODY, bold_color=INK,
                   space_after=t["bullet_gap"], line=t["body_line"])

    # ---- main: earlier experience
    main_head(main, "EARLIER EXPERIENCE", t["main_head"], t["main_gap"])
    for lead, rest in c["earlier"]:
        p = para(main, space_after=t["bullet_gap"], line=t["body_line"])
        text(p, lead, t["body"], INK, bold=True)
        text(p, rest, t["body"], BODY)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    variant = sys.argv[1]
    out = sys.argv[2]
    build(variant, out, content.DIALS[variant])
    print("wrote", out)
